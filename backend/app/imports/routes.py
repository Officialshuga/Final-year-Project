# ============================================================================
# BLUEPRINT: imports  (prefix: /api/imports)  — ADMIN ONLY
#   GET  /template          -> download a CSV template
#   POST /preview           -> upload CSV/XLSX, validate, return preview (no save)
#   POST /commit            -> save ONLY valid rows + create logins + temp pwds
# ----------------------------------------------------------------------------
# SECURITY:
#   - File size + type limited.
#   - Each imported staff gets a login with a RANDOM temp password + the
#     must_change_password flag. Temp passwords are hashed (bcrypt) and only
#     returned ONCE in the commit response for the admin to download.
#   - The import is written to the audit log.
# ============================================================================
import io
import csv
import re
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file, current_app
from flask_jwt_extended import jwt_required

from app.extensions import db
from app.models import Staff, Department, Shift, User
from app.utils.security import admin_required
from app.utils.audit import record_audit
from app.utils.password_gen import generate_temp_password
from datetime import timedelta

imports_bp = Blueprint("imports", __name__, url_prefix="/api/imports")

MAX_BYTES = 2 * 1024 * 1024  # 2 MB limit
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
COLUMNS = ["staff_id", "full_name", "email", "phone", "department", "shift"]


# ----------------------------------------------------------------------------
# TEMPLATE download
# ----------------------------------------------------------------------------
@imports_bp.get("/template")
@jwt_required()
@admin_required
def template():
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    writer.writerow(["GVU/2026/101", "Example Staff", "example@gvu.edu", "08030000000", "Human Resources", "Morning Shift"])
    data = io.BytesIO(buf.getvalue().encode())
    return send_file(data, as_attachment=True, download_name="staff_import_template.csv",
                     mimetype="text/csv")


# ----------------------------------------------------------------------------
# Parse an uploaded file into a list of dict rows.
# ----------------------------------------------------------------------------
def _read_rows(file_storage):
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if len(raw) > MAX_BYTES:
        raise ValueError("File too large (max 2 MB).")

    rows = []
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for r in reader:
            rows.append({k: (r.get(k) or "").strip() for k in COLUMNS})
    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip().lower() if c else "" for c in row]
                continue
            values = {header[i]: (str(row[i]).strip() if i < len(row) and row[i] is not None else "")
                      for i in range(len(header))}
            rows.append({k: values.get(k, "") for k in COLUMNS})
    else:
        raise ValueError("Unsupported file type. Upload .csv or .xlsx.")
    return rows


# ----------------------------------------------------------------------------
# Validate rows -> returns (valid_rows, error_rows) with reasons.
# ----------------------------------------------------------------------------
def _validate(rows):
    valid, errors = [], []
    seen_ids, seen_emails = set(), set()

    # Lookup maps (case-insensitive) for departments & shifts by name.
    depts = {d.name.lower(): d for d in Department.query.all()}
    shifts = {s.shift_name.lower(): s for s in Shift.query.all()}

    for i, r in enumerate(rows, start=2):  # row 1 is the header
        issues = []
        sid = r["staff_id"].strip()
        name = r["full_name"].strip()
        email = r["email"].strip().lower()
        dept_name = r["department"].strip()
        shift_name = r["shift"].strip()

        if not sid:
            issues.append("Missing staff ID")
        if len(name) < 3:
            issues.append("Full name required")
        if not EMAIL_RE.match(email):
            issues.append("Invalid email")

        # Duplicates within the file
        if sid and sid.lower() in seen_ids:
            issues.append("Duplicate staff ID in file")
        if email and email in seen_emails:
            issues.append("Duplicate email in file")

        # Duplicates against the database
        if sid and Staff.query.filter_by(staff_number=sid).first():
            issues.append("Staff ID already exists")
        if email and Staff.query.filter(db.func.lower(Staff.email) == email).first():
            issues.append("Email already exists")

        # Department / shift must exist (by name) if provided
        dept = depts.get(dept_name.lower()) if dept_name else None
        shift = shifts.get(shift_name.lower()) if shift_name else None
        if dept_name and not dept:
            issues.append(f"Unknown department '{dept_name}'")
        if shift_name and not shift:
            issues.append(f"Unknown shift '{shift_name}'")

        record = {
            "row": i, "staff_id": sid, "full_name": name, "email": email,
            "phone": r["phone"].strip(), "department": dept_name, "shift": shift_name,
            "department_id": dept.department_id if dept else None,
            "shift_id": shift.shift_id if shift else None,
        }
        if issues:
            record["errors"] = issues
            errors.append(record)
        else:
            seen_ids.add(sid.lower())
            seen_emails.add(email)
            valid.append(record)

    return valid, errors


# ----------------------------------------------------------------------------
# PREVIEW — validate only, no save.
# ----------------------------------------------------------------------------
@imports_bp.post("/preview")
@jwt_required()
@admin_required
def preview():
    if "file" not in request.files:
        return jsonify({"error": "Validation", "message": "No file uploaded (field 'file')"}), 400
    try:
        rows = _read_rows(request.files["file"])
    except ValueError as e:
        return jsonify({"error": "Validation", "message": str(e)}), 400

    valid, errors = _validate(rows)
    return jsonify({
        "total": len(rows), "valid_count": len(valid), "error_count": len(errors),
        "valid": valid, "errors": errors,
    }), 200


# ----------------------------------------------------------------------------
# COMMIT — save valid rows + create logins + return temp passwords ONCE.
# ----------------------------------------------------------------------------
@imports_bp.post("/commit")
@jwt_required()
@admin_required
def commit():
    if "file" not in request.files:
        return jsonify({"error": "Validation", "message": "No file uploaded (field 'file')"}), 400
    try:
        rows = _read_rows(request.files["file"])
    except ValueError as e:
        return jsonify({"error": "Validation", "message": str(e)}), 400

    valid, errors = _validate(rows)
    credentials = []  # shown once to the admin

    for r in valid:
        s = Staff(
            staff_number=r["staff_id"], full_name=r["full_name"], email=r["email"],
            phone=r["phone"] or None, department_id=r["department_id"], shift_id=r["shift_id"],
            status="Active", date_registered=datetime.utcnow(),
        )
        db.session.add(s)
        db.session.flush()

        # Username from email local-part; make unique.
        base = re.sub(r"[^a-z0-9._]", "", r["email"].split("@")[0].lower()) or "staff"
        username = base
        n = 1
        while User.query.filter_by(username=username).first():
            n += 1
            username = f"{base}{n}"

        temp = generate_temp_password(12)
        u = User(username=username, role="staff", staff_id=s.staff_id, is_active=True)
        u.set_password(temp)                     # bcrypt hash; plain never stored
        u.must_change_password = True
        u.temp_password_expires = datetime.utcnow() + timedelta(hours=24)
        db.session.add(u)

        credentials.append({"staff_id": r["staff_id"], "full_name": r["full_name"],
                            "username": username, "temp_password": temp})

    db.session.commit()
    record_audit("IMPORT_STAFF", "staff",
                 f"Imported {len(valid)} staff, skipped {len(errors)}")

    return jsonify({
        "imported": len(valid),
        "skipped": len(errors),
        "errors": errors,
        # Returned ONCE. The admin downloads this list; we never store plain pwds.
        "credentials": credentials,
        "note": "Temporary passwords are shown once. Each staff must change theirs on first login.",
    }), 201
